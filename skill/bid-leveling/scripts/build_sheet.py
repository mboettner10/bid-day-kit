#!/usr/bin/env python3
"""Build a bid leveling workbook from bids.json.

Usage: python3 build_sheet.py bids.json "Output.xlsx"
"""
import csv
import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

STATUS_LABEL = {
    "included": "Included",
    "excluded": "Excluded",
    "qualified": "Qualified",
    "not_mentioned": "Not mentioned",
}
STATUS_FILL = {
    "Included": "DCEFE3",
    "Excluded": "F6D5D5",
    "Qualified": "FBEBCB",
    "Not mentioned": "EDEDED",
}
HEADER_FILL = PatternFill("solid", fgColor="1F3B4D")
HEADER_FONT = Font(bold=True, color="FFFFFF")
CATEGORY_FILL = PatternFill("solid", fgColor="D9E2E8")
THIN = Side(style="thin", color="BFC7CC")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
MONEY = '"$"#,##0;[Red]-"$"#,##0'


def load_checklist(name):
    path = Path(__file__).resolve().parent.parent / "checklists" / f"{name}.csv"
    with open(path, newline="") as f:
        return list(csv.DictReader(f))


def _yn(first, last):
    from openpyxl.worksheet.datavalidation import DataValidation
    dv = DataValidation(type="list", formula1='"Y,N"', allow_blank=False)
    dv.add(f"C{first}:C{last}")
    return dv


def header(ws, row, values):
    for col, value in enumerate(values, 1):
        cell = ws.cell(row=row, column=col, value=value)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = BORDER


def build(data, out_path):
    checklist = load_checklist(data.get("checklist", "electrical"))
    bidders = data["bidders"]
    if len(bidders) < 2:
        sys.exit("Need at least two bidders to level.")

    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"
    level = wb.create_sheet("Leveling")
    questions = wb.create_sheet("Questions")
    extras = wb.create_sheet("Alternates & allowances")

    # Leveling sheet: item rows, two columns per bidder (status, adjustment $).
    title = f'{data.get("project", "Project")} - {data.get("trade", "Electrical")} bid leveling'
    level["A1"] = title
    level["A1"].font = Font(bold=True, size=14)
    level["A2"] = "Status comes from each quote's wording (hover a cell to see it). Set Applies to N for items not on this job. Enter your own cost in the Adjustment $ column for any gap you want to carry."
    level["A2"].font = Font(italic=True, color="555555")

    not_applicable = set(data.get("not_applicable", []))
    cols = ["ID", "Scope item", "Applies to this job?"]
    for b in bidders:
        cols += [b["name"], "Adjustment $"]
    header(level, 4, cols)
    level.row_dimensions[4].height = 34
    level.column_dimensions["A"].width = 6
    level.column_dimensions["B"].width = 48
    level.column_dimensions["C"].width = 11
    for i in range(len(bidders)):
        level.column_dimensions[get_column_letter(4 + 2 * i)].width = 18
        level.column_dimensions[get_column_letter(5 + 2 * i)].width = 14

    gap_counts = [0] * len(bidders)
    row = 5
    first_item_row = row
    current_category = None
    for item in checklist:
        if item["category"] != current_category:
            current_category = item["category"]
            cell = level.cell(row=row, column=1, value=current_category)
            for c in range(1, len(cols) + 1):
                level.cell(row=row, column=c).fill = CATEGORY_FILL
            cell.font = Font(bold=True)
            row += 1
        level.cell(row=row, column=1, value=item["id"]).border = BORDER
        level.cell(row=row, column=2, value=item["item"]).border = BORDER
        applies = "N" if item["id"] in not_applicable else "Y"
        level.cell(row=row, column=3, value=applies).border = BORDER
        for i, b in enumerate(bidders):
            entry = b.get("items", {}).get(item["id"], {})
            label = STATUS_LABEL.get(entry.get("status", "not_mentioned"), "Not mentioned")
            if label in ("Excluded", "Not mentioned") and applies == "Y":
                gap_counts[i] += 1
            cell = level.cell(row=row, column=4 + 2 * i, value=label)
            cell.fill = PatternFill("solid", fgColor=STATUS_FILL[label])
            cell.border = BORDER
            text = []
            if entry.get("quote"):
                page = f' (p. {entry["page"]})' if entry.get("page") else ""
                text.append(f'"{entry["quote"]}"{page}')
            if entry.get("note"):
                text.append(f'Note: {entry["note"]}')
            if text:
                cell.comment = Comment("\n".join(text), "Bid Day Kit", width=300, height=110)
            adj = level.cell(row=row, column=5 + 2 * i)
            adj.number_format = MONEY
            adj.border = BORDER
        row += 1
    last_item_row = row - 1

    # Unmatched quote items, so nothing in a quote is silently dropped.
    row += 1
    level.cell(row=row, column=1, value="In the quote but not on the checklist").font = Font(bold=True)
    row += 1
    for i, b in enumerate(bidders):
        for text in b.get("unmatched", []):
            level.cell(row=row, column=2, value=f'{b["name"]}: {text}')
            level.cell(row=row, column=2).alignment = Alignment(wrap_text=True)
            row += 1
    level.data_validations.dataValidation.append(_yn(first_item_row, last_item_row))
    level.freeze_panes = "D5"
    for ws in (level, summary):
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True

    # Summary sheet with live formulas.
    summary["A1"] = title
    summary["A1"].font = Font(bold=True, size=14)
    if data.get("bid_date"):
        summary["A2"] = f'Bids due {data["bid_date"]}'
    header(summary, 4, ["Bidder", "Quote date", "Valid (days)", "Base price", "Your adjustments", "Adjusted total", "Excluded", "Not mentioned", "Qualified"])
    for i, b in enumerate(bidders):
        r = 5 + i
        adj_col = get_column_letter(5 + 2 * i)
        status_col = get_column_letter(4 + 2 * i)
        rng = f"Leveling!{status_col}{first_item_row}:{status_col}{last_item_row}"
        app = f"Leveling!$C${first_item_row}:$C${last_item_row}"
        values = [
            b["name"],
            b.get("quote_date"),
            b.get("valid_days"),
            b.get("base_price"),
            f"=SUM(Leveling!{adj_col}{first_item_row}:{adj_col}{last_item_row})",
            f"=IF(D{r}=\"\",\"\",D{r}+E{r})",
            f'=COUNTIFS({rng},"Excluded",{app},"Y")',
            f'=COUNTIFS({rng},"Not mentioned",{app},"Y")',
            f'=COUNTIFS({rng},"Qualified",{app},"Y")',
        ]
        for c, v in enumerate(values, 1):
            cell = summary.cell(row=r, column=c, value=v)
            cell.border = BORDER
            if c in (4, 5, 6):
                cell.number_format = MONEY
        if b.get("base_price") is None:
            summary.cell(row=r, column=4).comment = Comment("No total price found in this quote.", "Bid Day Kit")
    last = 4 + len(bidders)
    summary.conditional_formatting.add(
        f"F5:F{last}",
        FormulaRule(formula=[f"AND(F5<>\"\",F5=MIN($F$5:$F${last}))"], fill=PatternFill("solid", fgColor="DCEFE3")),
    )
    summary.cell(row=last + 2, column=1, value="Adjusted total = base price + your adjustments for gaps. The lowest adjusted total is shaded; it is a comparison aid, not a recommendation.").font = Font(italic=True, color="555555")
    for col, width in zip("ABCDEFGHI", (28, 12, 11, 14, 16, 15, 10, 14, 10)):
        summary.column_dimensions[col].width = width

    # Questions per bidder.
    header(questions, 1, ["Bidder", "Question"])
    questions.column_dimensions["A"].width = 28
    questions.column_dimensions["B"].width = 90
    r = 2
    for b in bidders:
        for q in b.get("questions", []):
            questions.cell(row=r, column=1, value=b["name"])
            questions.cell(row=r, column=2, value=q).alignment = Alignment(wrap_text=True)
            r += 1

    # Alternates and allowances.
    header(extras, 1, ["Bidder", "Type", "Description", "Amount"])
    for col, width in zip("ABCD", (28, 12, 60, 14)):
        extras.column_dimensions[col].width = width
    r = 2
    for b in bidders:
        for kind, key, amount in (("Alternate", "alternates", "price"), ("Allowance", "allowances", "amount")):
            for x in b.get(key, []):
                extras.append([b["name"], kind, x.get("description"), x.get(amount)])
                extras.cell(row=r, column=4).number_format = MONEY
                r += 1

    wb.save(out_path)
    return gap_counts


def main():
    if len(sys.argv) != 3:
        sys.exit(__doc__)
    with open(sys.argv[1]) as f:
        data = json.load(f)
    gaps = build(data, sys.argv[2])
    for b, g in zip(data["bidders"], gaps):
        print(f'{b["name"]}: base {b.get("base_price")}, {g} gaps')
    print(f"Wrote {sys.argv[2]}")


if __name__ == "__main__":
    main()
